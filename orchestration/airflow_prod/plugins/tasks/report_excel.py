# report_excel.py
# -----------------------------------------------------------------------------
# Generatory zalacznikow Excel do raportu utrzymaniowego PKP:
#   1) report_steps_{date}.xlsx        - kroki daily + czas w sekundach
#   2) report_schema_sizes_{date}.xlsx - zajetosc per schemat (tab per schemat),
#                                        najwieksze obiekty wg MB
#
# Pliki sa TYMCZASOWE: katalog jest czyszczony przed kazdym generowaniem.
# Lokalizacja: <root stream_mpk>/data/<ENV>/RAPORT, gdzie <ENV> = DEV/PROD
# wyprowadzone z PKP_ENV (nazwa instancji). Root i db.py znajdowane dynamicznie.
#
# Test reczny (venv + env-y PKP):  python report_excel.py <pipeline_run_id>
# -----------------------------------------------------------------------------

import os
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# -- Odnalezienie root repo (stream_mpk) + import ingestion/db.py --------------
def _find_repo_root() -> Path:
    """Idzie w gore od tego pliku, zwraca katalog zawierajacy 'ingestion'."""
    here = Path(__file__).resolve()
    for parent in here.parents:
        if (parent / "ingestion" / "db.py").exists():
            return parent
    raise RuntimeError("Nie znaleziono katalogu 'ingestion' idac w gore od report_excel.py")


_REPO_ROOT = _find_repo_root()
if str(_REPO_ROOT / "ingestion") not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT / "ingestion"))
import db as _db   # ingestion/db.py -> _db.get_connection()


# -- Srodowisko (DEV/PROD) i katalog wyjsciowy --------------------------------
def _resolve_env(raw: str) -> str:
    """PKP_ENV to nazwa instancji -> wyprowadzamy segment DEV/PROD dla sciezki."""
    u = (raw or "DEV").strip().upper()
    if "PROD" in u:
        return "PROD"
    if "DEV" in u:
        return "DEV"
    return u   # gdy PKP_ENV jest juz czystym DEV/PROD lub czyms wlasnym


ENV = _resolve_env(os.environ.get("PKP_ENV", "DEV"))
REPORTS_DIR = str(_REPO_ROOT / "data" / ENV / "RAPORT")


# -- Konfiguracja raportu -----------------------------------------------------
# Schematy uwzgledniane w raporcie zajetosci. UZUPELNIJ/POPRAW pod swoja baze.
REPORT_SCHEMAS = ["SILVER", "GOLD", "MAINTENANCE", "STG"]

# Ile najwiekszych obiektow pokazac per schemat (None = wszystkie).
TOP_OBJECTS_PER_SCHEMA = None

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TOTAL_FONT = Font(bold=True)


# -- Pomocnicze ---------------------------------------------------------------
def cleanup_reports(pattern: str) -> None:
    """Kasuje z katalogu RAPORT tylko pliki pasujace do wzorca (np. 'report_steps_*.xlsx')."""
    p = Path(REPORTS_DIR)
    p.mkdir(parents=True, exist_ok=True)
    removed = 0
    for f in p.glob(pattern):
        if f.is_file():
            f.unlink()
            removed += 1
    print(f"[cleanup] Usunieto {removed} plikow ({pattern}) z {REPORTS_DIR}")


def _strip_timezone(value):
    """openpyxl nie zapisze datetime z tzinfo - zdejmujemy strefe przed zapisem."""
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


def _safe_fetch(conn, sql, params=None):
    """Zwraca (kolumny, wiersze) i nigdy nie wywala taska - pusty wynik = ([],[])."""
    try:
        with conn.cursor() as cur:
            cur.execute(sql, params or {})
            cols = [d[0] for d in cur.description] if cur.description else []
            rows = cur.fetchall()
        return cols, rows
    except Exception as e:
        print(f"[report_excel] Blad zapytania: {e}")
        return [], []


def _style_header(ws, columns, row=1):
    for col_idx, name in enumerate(columns, 1):
        c = ws.cell(row=row, column=col_idx, value=name)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(horizontal="center")


def _auto_width(ws):
    for col in ws.columns:
        longest = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(longest + 4, 60)


# -- 1) Excel: kroki daily + sekundy ------------------------------------------
def generate_steps_excel(run_date: str, pipeline_run_id: int, conn) -> str:
    """
    report_steps_{date}.xlsx - jeden arkusz: kroki runu z czasem w sekundach.
    Czas liczony w Pythonie z start_time/end_time (zachowuje ulamki sekund).
    """
    cols, rows = _safe_fetch(
        conn,
        """
        SELECT step_name, status, start_time, end_time
        FROM   maintenance.pipeline_run_step
        WHERE  pipeline_run_id = :run_id
          AND  step_name <> 'send_report'
        ORDER  BY start_time
        """,
        {"run_id": pipeline_run_id},
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "pipeline_steps"

    _style_header(ws, ["Krok", "Status", "Start", "Koniec", "Czas [s]"])

    total_seconds = 0.0
    r = 2
    for row in rows:
        step_name, status, start_t, end_t = row
        duration = None
        if start_t and end_t:
            duration = round((end_t - start_t).total_seconds(), 1)
            total_seconds += duration

        ws.cell(row=r, column=1, value=step_name)
        ws.cell(row=r, column=2, value=status)
        ws.cell(row=r, column=3, value=_strip_timezone(start_t))
        ws.cell(row=r, column=4, value=_strip_timezone(end_t))
        ws.cell(row=r, column=5, value=duration)
        r += 1

    ws.cell(row=r, column=1, value="RAZEM").font = TOTAL_FONT
    ws.cell(row=r, column=5, value=round(total_seconds, 1)).font = TOTAL_FONT

    ws.freeze_panes = "A2"
    _auto_width(ws)

    file_path = os.path.join(REPORTS_DIR, f"report_steps_{run_date}.xlsx")
    wb.save(file_path)
    print(f"[report_excel] Zapisano: {file_path} ({len(rows)} krokow)")
    return file_path


# -- 2) Excel: zajetosc per schemat (tab per schemat) -------------------------
def generate_schema_sizes_excel(run_date: str, conn, schemas=None) -> str:
    """
    report_schema_sizes_{date}.xlsx:
      - arkusz 'Podsumowanie' - suma MB per schemat + total,
      - arkusz per schemat    - obiekty posortowane malejaco wg MB.
    """
    schemas = schemas or REPORT_SCHEMAS

    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Podsumowanie"
    _style_header(summary, ["Schemat", "Rozmiar [MB]"])

    grand_total = 0.0
    summary_row = 2

    for schema in schemas:
        cols, rows = _safe_fetch(
            conn,
            """
            SELECT segment_name,
                   segment_type,
                   ROUND(SUM(bytes) / 1024 / 1024, 2) AS mb
            FROM   dba_segments
            WHERE  owner = :owner
            GROUP  BY segment_name, segment_type
            ORDER  BY SUM(bytes) DESC
            """,
            {"owner": schema},
        )

        schema_total = round(sum((row[2] or 0) for row in rows), 2)
        grand_total += schema_total

        summary.cell(row=summary_row, column=1, value=schema)
        summary.cell(row=summary_row, column=2, value=schema_total)
        summary_row += 1

        ws = wb.create_sheet(title=schema[:31])
        _style_header(ws, ["Obiekt", "Typ", "Rozmiar [MB]"])

        display_rows = rows if TOP_OBJECTS_PER_SCHEMA is None else rows[:TOP_OBJECTS_PER_SCHEMA]
        r = 2
        for seg_name, seg_type, mb in display_rows:
            ws.cell(row=r, column=1, value=seg_name)
            ws.cell(row=r, column=2, value=seg_type)
            ws.cell(row=r, column=3, value=mb)
            r += 1

        ws.cell(row=r, column=1, value="RAZEM").font = TOTAL_FONT
        ws.cell(row=r, column=3, value=schema_total).font = TOTAL_FONT

        ws.freeze_panes = "A2"
        _auto_width(ws)
        print(f"[report_excel] {schema}: {len(rows)} obiektow, {schema_total} MB")

    summary.cell(row=summary_row, column=1, value="TOTAL").font = TOTAL_FONT
    summary.cell(row=summary_row, column=2, value=round(grand_total, 2)).font = TOTAL_FONT
    summary.freeze_panes = "A2"
    _auto_width(summary)

    file_path = os.path.join(REPORTS_DIR, f"report_schema_sizes_{run_date}.xlsx")
    wb.save(file_path)
    print(f"[report_excel] Zapisano: {file_path} (total {round(grand_total, 2)} MB)")
    return file_path


# -- Test reczny --------------------------------------------------------------
if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uzycie: python report_excel.py <pipeline_run_id>")
        sys.exit(1)

    run_id = int(sys.argv[1])
    today = date.today().isoformat()

    print(f"[report_excel] ENV={ENV}  katalog={REPORTS_DIR}")
    cleanup_reports("report_steps_*.xlsx")
    cleanup_reports("report_schema_sizes_*.xlsx")
    conn = _db.get_connection()
    try:
        generate_steps_excel(today, run_id, conn)
        generate_schema_sizes_excel(today, conn)
    finally:
        conn.close()
    print("Gotowe.")